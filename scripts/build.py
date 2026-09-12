#!/usr/bin/env python3
"""Construit l'application à partir des données sources.

  python scripts/build.py                 -> génère app/data.js et met à jour la version de cache dans app/sw.js
  python scripts/build.py --xlsx          -> exporte aussi data/plan-comptable.xlsx (pour relecture et modifications dans Excel)
  python scripts/build.py --from-xlsx     -> réimporte data/plan-comptable.xlsx dans les JSON sources avant de construire
  python scripts/build.py --url https://… -> génère aussi le QR code app/icons/qr.png vers l'adresse de l'application
"""
import argparse, datetime, hashlib, json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
APP = os.path.join(ROOT, "app")

def load(name):
    with open(os.path.join(DATA, name), encoding="utf-8") as f:
        return json.load(f)

def save(name, obj):
    with open(os.path.join(DATA, name), "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)
        f.write("\n")

# ---------- Excel ----------
def export_xlsx(comptes, fiches, nature):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active; ws.title = "Comptes"
    ws.append(["Numéro", "Libellé", "Nature (vide = déduite de la classe)"])
    for c in comptes:
        ws.append([c["n"], c["l"], nature.get(c["n"], "")])
    ws2 = wb.create_sheet("Fiches")
    ws2.append(["Numéro", "En clair", "Mots-clés (séparés par ;)", "À ne pas confondre avec (numéro : raison, séparés par |)", "Comptes liés (séparés par ;)"])
    for c in comptes:
        f = fiches.get(c["n"])
        if not f: continue
        nc = " | ".join(f"{x[0]} : {x[1]}" if isinstance(x, list) else str(x) for x in f.get("nc", []))
        ws2.append([c["n"], f.get("i", ""), " ; ".join(f.get("k", [])), nc, " ; ".join(f.get("rel", []))])
    for sheet, widths in ((ws, (12, 90, 30)), (ws2, (12, 90, 50, 70, 30))):
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        for cell in sheet[1]:
            cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DCE3EA")
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    out = os.path.join(DATA, "plan-comptable.xlsx")
    wb.save(out)
    return out

def import_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(DATA, "plan-comptable.xlsx"), read_only=True)
    comptes, nature, fiches = [], {}, {}
    for n, l, nat in wb["Comptes"].iter_rows(min_row=2, values_only=True):
        if n is None: continue
        n = str(n).strip()
        comptes.append({"n": n, "l": (l or "").strip()})
        if nat: nature[n] = str(nat).strip()
    for row in wb["Fiches"].iter_rows(min_row=2, values_only=True):
        n, i, k, nc, rel = (list(row) + [None] * 5)[:5]
        if n is None: continue
        f = {}
        if i: f["i"] = str(i).strip()
        if k: f["k"] = [x.strip() for x in str(k).split(";") if x.strip()]
        if nc:
            f["nc"] = []
            for part in str(nc).split("|"):
                part = part.strip()
                if not part: continue
                m = re.match(r"(\d+)\s*:?\s*(.*)", part)
                if m: f["nc"].append([m.group(1), m.group(2).strip()])
        if rel: f["rel"] = [x.strip() for x in str(rel).split(";") if x.strip()]
        if f: fiches[str(n).strip()] = f
    comptes.sort(key=lambda c: c["n"])
    save("comptes.json", comptes); save("nature.json", nature); save("fiches.json", fiches)
    print(f"Excel importé : {len(comptes)} comptes, {len(fiches)} fiches")

# ---------- Contrôles ----------
def check(comptes, fiches, nature):
    nums = {c["n"] for c in comptes}
    problems = []
    if len(nums) != len(comptes): problems.append("numéros en double dans comptes.json")
    for n, f in fiches.items():
        if n not in nums: problems.append(f"fiche {n} : compte inexistant")
        for x in f.get("nc", []):
            m = x[0] if isinstance(x, list) else x
            if m not in nums: problems.append(f"fiche {n} : « ne pas confondre » vers {m} inexistant")
        for m in f.get("rel", []):
            if m not in nums: problems.append(f"fiche {n} : compte lié {m} inexistant")
    for n in nature:
        if n not in nums: problems.append(f"nature {n} : compte inexistant")
    return problems

# ---------- Construction ----------
def build(url=None):
    comptes, fiches, nature, classes = load("comptes.json"), load("fiches.json"), load("nature.json"), load("classes.json")
    problems = check(comptes, fiches, nature)
    for p in problems: print("ATTENTION :", p)
    version = datetime.date.today().strftime("%Y-%m-%d")
    data = {"version": version, "classes": classes, "accounts": comptes, "infos": fiches, "nature": nature}
    js = "window.PCG=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    with open(os.path.join(APP, "data.js"), "w", encoding="utf-8") as f:
        f.write(js)
    if url:
        import segno
        segno.make(url, error="m").save(os.path.join(APP, "icons", "qr.png"), scale=8, border=2, dark="#142230")
        with open(os.path.join(APP, "url.txt"), "w", encoding="utf-8") as f: f.write(url)
    # version de cache = empreinte de tous les fichiers servis (sauf sw.js lui-même)
    h = hashlib.sha256()
    for dirpath, _, files in os.walk(APP):
        for name in sorted(files):
            if name == "sw.js": continue
            with open(os.path.join(dirpath, name), "rb") as f: h.update(name.encode() + f.read())
    digest = h.hexdigest()[:10]
    sw_path = os.path.join(APP, "sw.js")
    with open(sw_path, encoding="utf-8") as f: sw = f.read()
    sw = re.sub(r'const VERSION = "[^"]*";', f'const VERSION = "{digest}";', sw)
    with open(sw_path, "w", encoding="utf-8") as f: f.write(sw)
    nb_i = sum(1 for f in fiches.values() if f.get("i"))
    print(f"data.js : {len(comptes)} comptes, {len(fiches)} fiches ({nb_i} explications), version {version}, cache {digest}, {len(js)//1024} Ko")
    return not problems

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", action="store_true"); ap.add_argument("--from-xlsx", action="store_true"); ap.add_argument("--url")
    a = ap.parse_args()
    if a.from_xlsx: import_xlsx()
    ok = build(a.url)
    if a.xlsx:
        print("Excel :", export_xlsx(load("comptes.json"), load("fiches.json"), load("nature.json")))
    sys.exit(0 if ok else 1)
